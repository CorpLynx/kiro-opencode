import io
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from ipredact import Redactor  # noqa: E402

try:
    import lxml  # noqa: F401
    from ipredact.office import redact_package, redact_xml
except ImportError:  # pragma: no cover
    lxml = None

KEY = "test-key-0123456789"


def requires(*mods):
    missing = []
    for m in mods:
        try:
            __import__(m)
        except ImportError:
            missing.append(m)
    return unittest.skipIf(missing, f"needs {', '.join(missing)}")


def run_cli(*args, check=True):
    return subprocess.run([sys.executable, "-m", "ipredact", *args], capture_output=True,
                          check=check, cwd=ROOT)


def zip_text(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


@requires("lxml")
class XmlPartTest(unittest.TestCase):
    def setUp(self):
        self.r = Redactor(KEY)
        self.addr = self.r.redact_text("10.1.2.3")

    def test_word_split_runs(self):
        xml = (b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
               b'<w:body><w:p><w:r><w:t>gw 10.1.</w:t></w:r><w:r><w:rPr><w:b/></w:rPr>'
               b'<w:t>2.3</w:t></w:r><w:r><w:t xml:space="preserve"> up</w:t></w:r></w:p></w:body></w:document>')
        out = redact_xml(self.r, xml, False).decode()
        self.assertNotIn("10.1.", out)
        self.assertIn(f"<w:t>gw {self.addr}</w:t>", out)
        self.assertIn("<w:b/>", out)  # formatting untouched

    def test_split_leaves_preserved_spaces(self):
        xml = (b'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
               b'<w:r><w:t>10.1</w:t></w:r><w:r><w:t>.2.3 is the gw</w:t></w:r></w:p>')
        out = redact_xml(self.r, xml, False).decode()
        self.assertIn(f"<w:t>{self.addr}</w:t>", out)
        self.assertIn('<w:t xml:space="preserve"> is the gw</w:t>', out)

    def test_numbers_in_drawings_dont_join_text(self):
        xml = (b'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
               b'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
               b'<w:r><w:t>v 1.2.3.</w:t></w:r><w:r><w:drawing><wp:posOffset>4</wp:posOffset>'
               b'</w:drawing></w:r></w:p>')
        self.assertIsNone(redact_xml(self.r, xml, False))

    def test_visio_mixed_content_and_shape_data(self):
        xml = (b'<PageContents xmlns="http://schemas.microsoft.com/office/visio/2012/main"><Shapes>'
               b'<Shape><Section N="Property"><Row N="IP"><Cell N="Value" V="10.1.2.3"/></Row></Section>'
               b'<Text><cp IX="0"/>core-rtr 10.1.<cp IX="1"/>2.3</Text></Shape></Shapes></PageContents>')
        out = redact_xml(self.r, xml, False).decode()
        self.assertNotIn("10.1.", out)
        self.assertIn(f'V="{self.addr}"', out)
        self.assertIn(f"core-rtr {self.addr}", out)

    def test_odf_spans(self):
        xml = (b'<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
               b'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><text:p>'
               b'peer <text:span>10.1</text:span>.2.3</text:p></office:body></office:document-content>')
        out = redact_xml(self.r, xml, False).decode()
        self.assertIn(f"peer <text:span>{self.addr}</text:span></text:p>", out)

    def test_unchanged_returns_none_and_entities_not_expanded(self):
        self.assertIsNone(redact_xml(self.r, b"<a>nothing here</a>", False))
        xxe = (b'<?xml version="1.0"?><!DOCTYPE a [<!ENTITY x SYSTEM "file:///etc/passwd">]>'
               b'<a>&x; 10.1.2.3</a>')
        out = redact_xml(self.r, xxe, False)
        self.assertNotIn(b"root:", out)

    def test_round_trip(self):
        xml = (b'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
               b'<w:r><w:t>10.1</w:t></w:r><w:r><w:t>.2.3</w:t></w:r></w:p>')
        red = redact_xml(self.r, xml, False)
        back = redact_xml(Redactor(KEY), red, True)
        self.assertIn(b"10.1.2.3", back)


@requires("lxml", "docx")
class DocxTest(unittest.TestCase):
    def test_docx_everywhere(self):
        import docx
        d = docx.Document()
        p = d.add_paragraph("core gateway ")
        p.add_run("10.20.").italic = True
        p.add_run("30.1")
        d.sections[0].header.paragraphs[0].text = "site 10.20.0.0/16"
        table = d.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "2001:db8:abcd:12::1"
        d.add_comment(p.runs, text="ask about 10.20.30.1", author="me")
        # A hyperlink pointing at an internal address.
        rid = d.part.relate_to("http://10.20.30.1/admin",
                               "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                               is_external=True)
        buf = io.BytesIO()
        d.save(buf)
        data = buf.getvalue()
        # Add a thumbnail the way Word does.
        data = add_thumbnail(data)

        r = Redactor(KEY)
        out = redact_package(r, data, False, notes := [])
        parts = zip_text(out)
        blob = b"".join(parts.values())
        for leak in (b"10.20.", b"2001:db8:abcd"):
            self.assertNotIn(leak, blob)
        self.assertNotIn("docProps/thumbnail.jpeg", parts)
        self.assertNotIn(b"thumbnail", parts["_rels/.rels"])
        self.assertTrue(any("thumbnail" in n for n in notes))

        red = docx.Document(io.BytesIO(out))  # still a valid document
        self.assertEqual(red.paragraphs[0].text, "core gateway " + r.redact_text("10.20.30.1"))
        self.assertTrue(red.paragraphs[0].runs[1].italic)
        self.assertIn(r.redact_text("http://10.20.30.1/admin"), parts["word/_rels/document.xml.rels"].decode())
        self.assertTrue(rid)

        back = redact_package(Redactor(KEY), out, True, [])
        self.assertEqual(docx.Document(io.BytesIO(back)).paragraphs[0].text, "core gateway 10.20.30.1")


def add_thumbnail(data):
    """Make sure the package has a thumbnail the way Word saves one (the
    python-docx template already ships with one)."""
    src = zip_text(data)
    rels = src["_rels/.rels"]
    if b"thumbnail" not in rels:
        rels = rels.replace(
            b"</Relationships>",
            b'<Relationship Id="rIdThumb" Type="http://schemas.openxmlformats.org/package/2006/'
            b'relationships/metadata/thumbnail" Target="docProps/thumbnail.jpeg"/></Relationships>')
    src["_rels/.rels"] = rels
    src["docProps/thumbnail.jpeg"] = b"\xff\xd8\xff fake jpeg showing 10.20.30.1"
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for name, content in src.items():
            z.writestr(name, content)
    return out.getvalue()


@requires("lxml", "pptx")
class PptxTest(unittest.TestCase):
    def test_slide_and_notes(self):
        import pptx
        from pptx.util import Inches
        prs = pptx.Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "DC 10.40.0.0/16"
        tf = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame
        para = tf.paragraphs[0]
        para.add_run().text = "fw 10.40."
        para.add_run().text = "1.1"
        slide.notes_slide.notes_text_frame.text = "remember 10.40.1.1"
        buf = io.BytesIO()
        prs.save(buf)
        r = Redactor(KEY)
        out = redact_package(r, buf.getvalue(), False, [])
        self.assertNotIn(b"10.40.", b"".join(zip_text(out).values()))
        red = pptx.Presentation(io.BytesIO(out)).slides[0]
        texts = [s.text_frame.text for s in red.shapes if s.has_text_frame]
        self.assertIn("fw " + r.redact_text("10.40.1.1"), texts)


@requires("lxml", "openpyxl")
class XlsxTest(unittest.TestCase):
    def test_charts_formulas_comments_survive(self):
        import openpyxl
        from openpyxl.chart import BarChart, Reference
        from openpyxl.comments import Comment
        with tempfile.TemporaryDirectory() as d:
            src, dst = os.path.join(d, "in.xlsx"), os.path.join(d, "out.xlsx")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "10.1.0.0"
            ws.append(["subnet", "hosts"])
            ws.append(["10.1.1.0/24", 50])
            ws.append(["10.1.2.0/24", 70])
            ws["C1"] = '="gw "&"10.1.1.1"'
            ws["A2"].comment = Comment("owned by 10.1.9.9", "me")
            ws["A3"].hyperlink = "https://10.1.2.2/ui"
            chart = BarChart()
            chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=3))
            ws.add_chart(chart, "E2")
            wb.save(src)

            run_cli(src, "-o", dst, "-k", KEY)
            with open(dst, "rb") as f:
                parts = zip_text(f.read())
            self.assertTrue(any(n.startswith("xl/charts/") for n in parts))  # openpyxl would drop this
            self.assertNotIn(b"10.1.", b"".join(parts.values()))

            r = Redactor(KEY)
            ws = openpyxl.load_workbook(dst).active
            self.assertEqual(ws.title, r.redact_text("10.1.0.0"))
            self.assertEqual(ws["A2"].value, r.redact_text("10.1.1.0/24"))
            self.assertEqual(ws["B2"].value, 50)
            self.assertEqual(ws["C1"].value, '="gw "&"' + r.redact_text("10.1.1.1") + '"')
            self.assertNotIn("10.1.", ws["A2"].comment.text)
            self.assertNotIn("10.1.", ws["A3"].hyperlink.target)

            back = os.path.join(d, "back.xlsx")
            run_cli(dst, "-o", back, "-k", KEY, "-r")
            self.assertEqual(openpyxl.load_workbook(back).active["A2"].value, "10.1.1.0/24")


@requires("lxml")
class PackageTest(unittest.TestCase):
    def make(self, files):
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as z:
            for name, content in files.items():
                z.writestr(name, content)
        return out.getvalue()

    def test_embedded_package_and_opaque_objects(self):
        inner = self.make({"[Content_Types].xml": b"<Types/>",
                           "xl/sharedStrings.xml": b'<sst xmlns="http://schemas.openxmlformats.org/'
                                                   b'spreadsheetml/2006/main"><si><t>10.9.9.9</t></si></sst>'})
        outer = self.make({"[Content_Types].xml": b"<Types/>",
                           "word/embeddings/Microsoft_Excel_Worksheet.xlsx": inner,
                           "word/embeddings/oleObject1.bin": b"\xd0\xcf\x11\xe0 10.9.9.9"})
        notes = []
        out = redact_package(Redactor(KEY), outer, False, notes)
        embedded = zip_text(zip_text(out)["word/embeddings/Microsoft_Excel_Worksheet.xlsx"])
        self.assertNotIn(b"10.9.9.9", embedded["xl/sharedStrings.xml"])
        self.assertTrue(any("not inspected" in n for n in notes))

    def test_odf_mimetype_stays_first_and_stored(self):
        files = {"mimetype": b"application/vnd.oasis.opendocument.text",
                 "content.xml": b'<x xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
                                b'<text:p>10.9.9.9</text:p></x>'}
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as z:
            z.writestr(zipfile.ZipInfo("mimetype"), files["mimetype"])  # stored
            z.writestr("content.xml", files["content.xml"], zipfile.ZIP_DEFLATED)
        red = redact_package(Redactor(KEY), out.getvalue(), False, [])
        with zipfile.ZipFile(io.BytesIO(red)) as z:
            first = z.infolist()[0]
            self.assertEqual(first.filename, "mimetype")
            self.assertEqual(first.compress_type, zipfile.ZIP_STORED)
            self.assertNotIn(b"10.9.9.9", z.read("content.xml"))

    def test_plain_zip_is_skipped(self):
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "x.zip")
            with open(p, "wb") as f:
                f.write(self.make({"a.txt": b"10.1.2.3"}))
            res = run_cli(p, "-o", os.path.join(d, "o.zip"), "-k", KEY, check=False)
            self.assertEqual(res.returncode, 1)
            self.assertIn(b"isn't an Office", res.stderr)


if __name__ == "__main__":
    unittest.main()

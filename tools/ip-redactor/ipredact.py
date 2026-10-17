#!/usr/bin/env python3
"""Structure-preserving IPv4 / IPv6 redaction.

IPv4: the first two octets are rewritten with different numbers and the last
two are kept. IPv6: the routing prefix (first 48 bits by default) is
rewritten and the subnet ID and interface ID are kept. The result keeps its
shape (same subnets, same hosts, same relationships) but no longer points at
real infrastructure.

The mapping is keyed by a secret and is:

  * Consistent  - the same input always becomes the same output for a key,
                  across lines, files and runs.
  * One-to-one  - two different addresses never collide.
  * Prefix-preserving - addresses that share a prefix (at /8 or longer for
                  IPv4, at any length for IPv6) still share it afterwards, so
                  CIDR containment survives redaction.
  * Always different - for IPv4 both of the first two octets change; for IPv6
                  every rewritten hextet changes.
  * Reversible  - with the key, --reverse restores the original addresses.

IPv4 first octets go through a keyed single-cycle permutation that never
produces 0, 127 or 255 (and 169.254/16 is stepped around); second octets (and the IPv6 prefix) go through a
keyed bit-flip tree in the style of Crypto-PAn. IPv6 keeps the top 3 bits
(and the top 10 inside e000::/3) so global unicast stays global unicast and
ULA stays ULA, and nothing is ever turned into link-local or multicast.
"""

import argparse
import csv
import hashlib
import hmac
import io
import ipaddress
import os
import re
import secrets
import sys
import tempfile

OCTET = r"(?:25[0-5]|2[0-4]\d|1\d\d|0?\d?\d)"
IPV4 = rf"{OCTET}\.{OCTET}\.{OCTET}\.{OCTET}"
HEX = r"[0-9A-Fa-f]{0,4}"

ARPA_PAT = rf"""
    (?<![\w.])(?:
        (?P<arpa4>(?:{OCTET}\.){{1,4}})(?P<sfx4>in-addr\.arpa)
      | (?P<arpa6>(?:[0-9A-Fa-f]\.){{1,32}})(?P<sfx6>ip6\.arpa)
    )(?![\w-])
"""

V6_PAT = rf"""
    (?<![\w:.])(?P<v6>(?:{HEX}:){{2,7}}(?:{IPV4}|{HEX}))
    (?P<zone>%[\w-]+)?
    (?:/(?P<len6>12[0-8]|1[01]\d|[1-9]?\d)(?!\d))?
    (?![\w:]|\.\w)
"""

# An IPv4 address, optionally followed by a CIDR length or by a dotted mask
# ("10.0.0.0 255.0.0.0", "network 10.0.0.0 mask 255.0.0.0",
# "permit 10.0.0.0 0.255.255.255"). The mask decides whether a network
# address needs re-aligning after its prefix bits are rewritten.
V4_PAT = rf"""
    (?<![\d.])(?P<v4>{IPV4})
    (?: /(?P<len4>3[0-2]|[12]?\d)(?!\d)
      | (?P<sep>[ \t]+(?:mask[ \t]+)?)(?P<mask>{IPV4}) )?
    (?!\.?\d)
"""

TOKEN_RE = re.compile(f"{ARPA_PAT}|{V6_PAT}|{V4_PAT}", re.VERBOSE | re.IGNORECASE)
V4_RE = re.compile(V4_PAT, re.VERBOSE | re.IGNORECASE)

# AWS-style hostnames embed the address with dashes: ip-10-1-2-3.ec2.internal
DASHED_RE = re.compile(
    rf"(?<=ip-)(?P<ip>{OCTET}-{OCTET}-{OCTET}-{OCTET})(?![\d-]*\d)",
    re.IGNORECASE,
)

DEFAULT_EXCLUDES = (
    "0.0.0.0/8",        # "this network", wildcard masks like 0.0.0.255
    "127.0.0.0/8",      # loopback
    "169.254.0.0/16",   # link-local, incl. the 169.254.169.254 metadata IP
    "255.0.0.0/8",      # netmasks and broadcast
    "::/127",           # unspecified and loopback
    "fe80::/10",        # link-local
    "ff00::/8",         # multicast
)

# First octets the IPv4 permutation works over. 0, 127 and 255 are left out
# so a redacted address never lands in them.
V4_FIRST_OCTETS = tuple(o for o in range(256) if o not in (0, 127, 255))

# IPv6 prefixes whose low 32 bits are an embedded IPv4 address.
V6_EMBEDS_V4 = (
    int(ipaddress.IPv6Address("::ffff:0:0")) >> 32,   # IPv4-mapped
    int(ipaddress.IPv6Address("64:ff9b::")) >> 32,    # NAT64 well-known prefix
)

# 169.254/16 as a 16-bit IPv4 prefix. The first-octet permutation keeps
# outputs out of 0/8, 127/8 and 255/8; this one /16 is avoided separately.
LINK_LOCAL16 = (169 << 8) | 254

V6_BITS_CHOICES = (16, 32, 48, 64)
KEY_ENV = "IPREDACT_KEY"
MIN_KEY_LEN = 16
XLSX_EXTS = (".xlsx", ".xlsm")
UTF16_BOMS = ((b"\xff\xfe", "utf-16-le"), (b"\xfe\xff", "utf-16-be"))


def mask_len(value):
    """Prefix length of a contiguous netmask or wildcard mask, else None."""
    if value == 0:
        return 32  # 0.0.0.0 next to an address is a host wildcard
    if value >> 31:
        inverted = ~value & 0xFFFFFFFF
        if inverted & (inverted + 1) == 0:
            return 32 - inverted.bit_length()
        return None
    if value & (value + 1) == 0:
        return 32 - value.bit_length()
    return None


def fmt_v4(value, template):
    """Format an IPv4 int; zero-padded input (010.001.002.003) stays padded."""
    padded = any(len(o) > 1 and o.startswith("0") for o in template.split("."))
    octets = [str((value >> (24 - 8 * i)) & 0xFF) for i in range(4)]
    return ".".join(o.zfill(3) for o in octets) if padded else ".".join(octets)


def fmt_v6(value, template):
    """Format an IPv6 int in the template's style: compressed or not,
    zero-padded or not, upper or lower case, dotted IPv4 tail or not."""
    dotted = "." in template
    compressed = "::" in template
    head = template.rsplit(":", 1)[0] if dotted else template
    padded = not compressed and all(len(g) == 4 for g in head.split(":"))
    hextets = [(value >> (112 - 16 * i)) & 0xFFFF for i in range(8)]
    if dotted:
        hextets = hextets[:6]
    items = [f"{h:04x}" if padded else f"{h:x}" for h in hextets]
    if dotted:
        items.append(str(ipaddress.IPv4Address(value & 0xFFFFFFFF)))
    text = ":".join(items)
    if compressed:
        # RFC 5952: collapse the longest run (2+) of zero hextets, leftmost wins.
        best_start, best_len, start = None, 1, None
        for i, h in enumerate(hextets + [1]):
            if h == 0 and start is None:
                start = i
            elif h != 0 and start is not None:
                if i - start > best_len:
                    best_start, best_len = start, i - start
                start = None
        if best_start is not None:
            text = ":".join(items[:best_start]) + "::" + ":".join(items[best_start + best_len:])
    return text.upper() if any(c in "ABCDEF" for c in template) else text


class Redactor:
    def __init__(self, key, excludes=DEFAULT_EXCLUDES, dashed=True, v6_bits=48):
        if isinstance(key, str):
            key = key.encode()
        if not key:
            raise ValueError("key must not be empty")
        if v6_bits not in V6_BITS_CHOICES:
            raise ValueError(f"v6_bits must be one of {V6_BITS_CHOICES}")
        self._key = key
        nets = [ipaddress.ip_network(n, strict=False) for n in excludes]
        self._excludes = {
            v: [(int(n.network_address), int(n.broadcast_address)) for n in nets if n.version == v]
            for v in (4, 6)
        }
        ll = ipaddress.ip_network("169.254.0.0/16")
        self._avoid_ll = any(n.version == 4 and ll.subnet_of(n) for n in nets)
        self._v4_fwd, self._v4_rev, self._v6_fwd, self._v6_rev = {}, {}, {}, {}
        self._text_cache = {}  # (text, reverse, prefix_len) -> formatted output
        self._v6_valid = {}    # IPv6 candidate text -> is it a real address
        self._dashed = dashed
        self.v6_bits = v6_bits
        self._flip_cache = {}
        shuffled = self._keyed_cycle("v4-octet1", V4_FIRST_OCTETS)
        self._oct1 = dict(zip(V4_FIRST_OCTETS, shuffled))
        self._oct1_inv = {v: k for k, v in self._oct1.items()}
        self.seen = {}           # original address -> redacted address
        self.collisions = set()  # redacted outputs that landed in an excluded range
        self.count = 0

    # -- keyed primitives -------------------------------------------------

    def _prf(self, *parts):
        msg = "|".join(str(p) for p in parts).encode()
        return hmac.new(self._key, msg, hashlib.sha256).digest()

    def _keyed_cycle(self, label, items):
        """Sattolo's shuffle: a keyed permutation that is a single cycle, so
        no item maps to itself."""
        items = list(items)
        for i in range(len(items) - 1, 0, -1):
            j = int.from_bytes(self._prf(label, i)[:8], "big") % i
            items[i], items[j] = items[j], items[i]
        return items

    def _flip(self, label, bit, context, prefix, rule):
        forced = rule(bit, prefix)
        if forced is not None:
            return forced
        cache_key = (label, bit, context, prefix)
        if cache_key not in self._flip_cache:
            self._flip_cache[cache_key] = self._prf(label, bit, context, prefix)[0] & 1
        return self._flip_cache[cache_key]

    def _tree(self, label, value, width, context, rule, reverse):
        """Prefix-preserving bijection on `width` bits: each bit is flipped by
        a keyed function of the original bits above it."""
        result = 0
        for bit in range(width):
            shift = width - 1 - bit
            orig_above = (result if reverse else value) >> (shift + 1)
            flip = self._flip(label, bit, context, orig_above, rule)
            result |= (((value >> shift) & 1) ^ flip) << shift
        return result

    @staticmethod
    def _v4_rule(bit, prefix):
        return 1 if bit == 0 else None  # always change the second octet

    def _v6_rule(self, bit, prefix):
        if bit < 3:
            return 0  # keep the address type (global unicast stays 2000::/3)
        if bit < 10 and prefix >> (bit - 3) == 0b111:
            return 0  # keep ULA / link-local / multicast distinct inside e000::/3
        if bit % 16 == 15:
            return 1  # always change every rewritten hextet
        return None

    # -- raw maps ---------------------------------------------------------

    def _v4_prefix(self, p, reverse):
        o1, o2 = p >> 8, p & 0xFF
        if reverse:
            o1 = self._oct1_inv.get(o1, o1)
            o2 = self._tree("v4-octet2", o2, 8, o1, self._v4_rule, True)
        else:
            o2 = self._tree("v4-octet2", o2, 8, o1, self._v4_rule, False)
            o1 = self._oct1.get(o1, o1)
        return (o1 << 8) | o2

    def _map_v4_int(self, value, reverse):
        top = value >> 16
        cache = self._v4_rev if reverse else self._v4_fwd
        if top not in cache:
            if reverse:
                # Undo the forward step below: the slot 169.254 would have
                # taken was handed to whichever prefix really mapped there.
                if self._avoid_ll and top == self._v4_prefix(LINK_LOCAL16, False):
                    out = self._v4_prefix(LINK_LOCAL16, True)
                else:
                    out = self._v4_prefix(top, True)
            else:
                out = self._v4_prefix(top, False)
                if self._avoid_ll and out == LINK_LOCAL16:
                    # 169.254/16 is excluded, so its own image is free: take
                    # it. One-to-one and reversible; only this one /16 per key
                    # loses prefix preservation.
                    out = self._v4_prefix(LINK_LOCAL16, False)
            cache[top] = out
        return (cache[top] << 16) | (value & 0xFFFF)

    def _map_v6_int(self, value, reverse):
        low_bits = 128 - self.v6_bits
        top = value >> low_bits
        cache = self._v6_rev if reverse else self._v6_fwd
        if top not in cache:
            cache[top] = self._tree("v6", top, self.v6_bits, 0, self._v6_rule, reverse)
        return (cache[top] << low_bits) | (value & ((1 << low_bits) - 1))

    # -- whole addresses --------------------------------------------------

    def is_excluded(self, ip):
        value = int(ip)
        return any(lo <= value <= hi for lo, hi in self._excludes[ip.version])

    def map_addr(self, ip, prefix_len=None, reverse=False):
        """Map an IPv4Address / IPv6Address.

        With prefix_len, an address that is the network address of that
        prefix stays a valid network address (re-aligned if the prefix ends
        inside the rewritten bits), and very short prefixes (< /8 IPv4,
        < /16 IPv6) are left alone since they identify nothing.
        """
        if self.is_excluded(ip):
            return ip
        value, width = int(ip), ip.max_prefixlen
        if ip.version == 6 and value >> 32 in V6_EMBEDS_V4:
            inner_len = None if prefix_len is None else max(prefix_len - 96, 0)
            inner = self.map_addr(ipaddress.IPv4Address(value & 0xFFFFFFFF), inner_len, reverse)
            return ipaddress.IPv6Address((value >> 32 << 32) | int(inner))

        shortest, rewritten = (8, 16) if ip.version == 4 else (16, self.v6_bits)
        aligned = False
        if prefix_len is not None:
            hostmask = (1 << (width - prefix_len)) - 1
            aligned = value & hostmask == 0
            if aligned and prefix_len < shortest:
                return ip
        out = (self._map_v4_int if ip.version == 4 else self._map_v6_int)(value, reverse)
        if aligned and prefix_len < rewritten:
            out &= ~((1 << (width - prefix_len)) - 1)
        out = ipaddress.IPv6Address(out) if ip.version == 6 else ipaddress.IPv4Address(out)
        self.count += 1
        if not reverse:
            self.seen[ip] = out
            if self.is_excluded(out):
                self.collisions.add(out)
        return out

    def map_ip(self, text, reverse=False, prefix_len=None):
        """Map an address given as text, keeping its formatting."""
        cache_key = (text, reverse, prefix_len)
        cached = self._text_cache.get(cache_key)
        if cached is not None:
            if cached != text:
                self.count += 1
            return cached
        ip = ipaddress.ip_address(text if ":" in text else ".".join(str(int(o)) for o in text.split(".")))
        out = self.map_addr(ip, prefix_len, reverse)
        if out == ip:
            result = text
        else:
            result = fmt_v4(int(out), text) if ip.version == 4 else fmt_v6(int(out), text)
        self._text_cache[cache_key] = result
        return result

    # -- text -------------------------------------------------------------

    def _v4_token(self, m, reverse):
        ip, plen, sep, mask = m.group("v4", "len4", "sep", "mask")
        if plen is not None:
            return f"{self.map_ip(ip, reverse, int(plen))}/{plen}"
        if mask is not None:
            mlen = mask_len(int(ipaddress.IPv4Address(".".join(str(int(o)) for o in mask.split(".")))))
            if mlen is None:
                # Not a mask, just two addresses in a row.
                return f"{self.map_ip(ip, reverse)}{sep}{self.map_ip(mask, reverse)}"
            return f"{self.map_ip(ip, reverse, mlen)}{sep}{mask}"
        return self.map_ip(ip, reverse)

    def _token(self, m, reverse):
        g = m.groupdict()
        if g.get("arpa4") is not None:
            octets = g["arpa4"].rstrip(".").split(".")
            n = len(octets)
            value = 0
            for i, o in enumerate(reversed(octets)):
                value |= int(o) << (24 - 8 * i)
            out = self.map_addr(ipaddress.IPv4Address(value), 8 * n if n < 4 else None, reverse)
            new = [str((int(out) >> (24 - 8 * i)) & 0xFF) for i in range(n)]
            return ".".join(reversed(new)) + "." + g["sfx4"]
        if g.get("arpa6") is not None:
            nibbles = g["arpa6"].rstrip(".").split(".")
            n = len(nibbles)
            value = int("".join(reversed(nibbles)), 16) << (128 - 4 * n)
            out = self.map_addr(ipaddress.IPv6Address(value), 4 * n if n < 32 else None, reverse)
            new = f"{int(out) >> (128 - 4 * n):0{n}x}"
            if g["arpa6"] != g["arpa6"].lower():
                new = new.upper()
            return ".".join(reversed(new)) + "." + g["sfx6"]
        if g.get("v6") is not None:
            text = g["v6"]
            valid = self._v6_valid.get(text)
            if valid is None:
                try:
                    ipaddress.IPv6Address(text)
                    valid = any(c.isdigit() for c in text)  # skip "dead::beef", "a::b"
                except ValueError:
                    valid = False
                self._v6_valid[text] = valid
            if not valid:
                # Not IPv6 (a MAC, a timestamp, ...). It may still carry an
                # IPv4 address, so don't let it shield one from redaction.
                return V4_RE.sub(lambda m4: self._v4_token(m4, reverse), m.group(0))
            plen = g["len6"]
            out = self.map_ip(text, reverse, int(plen) if plen is not None else None)
            return out + (g["zone"] or "") + (f"/{plen}" if plen is not None else "")
        return self._v4_token(m, reverse)

    def redact_text(self, text, reverse=False):
        text = TOKEN_RE.sub(lambda m: self._token(m, reverse), text)
        if self._dashed:
            text = DASHED_RE.sub(
                lambda m: self.map_ip(m.group("ip").replace("-", "."), reverse).replace(".", "-"),
                text,
            )
        return text


# -- file handling ------------------------------------------------------------

class SkipFile(Exception):
    pass


def current_umask():
    mask = os.umask(0)
    os.umask(mask)
    return mask


def write_bytes(path, data, mode=None):
    """Write atomically: a crash never leaves a half-written file behind."""
    if path == "-":
        sys.stdout.buffer.write(data)
        sys.stdout.flush()
        return
    path = os.path.realpath(path)  # replace a symlink's target, not the link
    if mode is None:
        try:
            mode = os.stat(path).st_mode & 0o7777
        except FileNotFoundError:
            mode = 0o666 & ~current_umask()
    fd, tmp = tempfile.mkstemp(dir=os.path.dirname(path), prefix=".ipredact-")
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(data)
        os.chmod(tmp, mode)
        os.replace(tmp, path)
    except BaseException:
        if os.path.exists(tmp):
            os.unlink(tmp)
        raise


def redact_bytes(redactor, data, reverse):
    for bom, enc in UTF16_BOMS:
        if data.startswith(bom):
            text = data[len(bom):].decode(enc, "surrogatepass")
            return bom + redactor.redact_text(text, reverse).encode(enc, "surrogatepass")
    if b"\x00" in data[:8192]:
        raise SkipFile("looks binary")
    text = data.decode("utf-8", "surrogateescape")
    return redactor.redact_text(text, reverse).encode("utf-8", "surrogateescape")


def redact_xlsx(redactor, src, reverse):
    try:
        import openpyxl
        from openpyxl.comments import Comment
    except ImportError:
        raise SkipFile(".xlsx support needs openpyxl (pip install openpyxl)")
    wb = openpyxl.load_workbook(src, keep_vba=src.lower().endswith(".xlsm"))
    for ws in wb.worksheets:
        title = redactor.redact_text(ws.title, reverse)
        if title != ws.title:
            if len(title) <= 31 and title not in wb.sheetnames:
                ws.title = title
            else:
                warn(f"{src}: could not rename sheet {ws.title!r}; its name still contains an address")
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = redactor.redact_text(cell.value, reverse)
                if cell.comment is not None:
                    cell.comment = Comment(redactor.redact_text(cell.comment.text, reverse),
                                           cell.comment.author)
                if cell.hyperlink is not None and cell.hyperlink.target:
                    cell.hyperlink.target = redactor.redact_text(cell.hyperlink.target, reverse)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def warn(msg):
    print(f"ipredact: {msg}", file=sys.stderr)


def collect_inputs(paths):
    """Expand directories into (source, path relative to the argument)."""
    out = []
    for p in paths:
        if p == "-":
            out.append(("-", "-"))
        elif os.path.isdir(p):
            for root, dirs, files in os.walk(p):
                dirs.sort()
                for name in sorted(files):
                    full = os.path.join(root, name)
                    out.append((full, os.path.relpath(full, p)))
        else:
            out.append((p, os.path.basename(p)))
    return out


def plan_outputs(inputs, args):
    plan, seen = [], {}
    for src, rel in inputs:
        if args.in_place:
            dst = src
        elif args.output:
            dst = args.output
        elif args.out_dir:
            dst = os.path.join(args.out_dir, rel)
        else:
            dst = "-"
        if dst != "-" and not args.in_place:
            real = os.path.realpath(dst)
            if src != "-" and os.path.exists(dst) and os.path.samefile(src, dst):
                sys.exit(f"error: {src} would overwrite itself; use --in-place for that")
            if real in seen:
                sys.exit(f"error: {seen[real]} and {src} would both write {dst}")
            seen[real] = src
        plan.append((src, dst))
    return plan


def resolve_key(args):
    if args.key:
        return args.key, False
    if args.key_file:
        with open(args.key_file) as f:
            key = f.read().strip()
        if not key:
            sys.exit(f"error: key file {args.key_file} is empty")
        return key, False
    if os.environ.get(KEY_ENV):
        return os.environ[KEY_ENV], False
    if args.reverse:
        sys.exit(f"error: --reverse needs the key used to redact (--key, --key-file or ${KEY_ENV})")
    return secrets.token_hex(16), True


def build_parser():
    p = argparse.ArgumentParser(
        description="Redact IPv4 and IPv6 addresses by rewriting their network "
                    "prefix while keeping subnet structure and host relationships intact.",
    )
    p.add_argument("files", nargs="*", default=["-"],
                   help="text, CSV or .xlsx/.xlsm files, or directories to walk (default: stdin)")
    out = p.add_mutually_exclusive_group()
    out.add_argument("-o", "--output", help="output file (single input only)")
    out.add_argument("-d", "--out-dir", help="write redacted files into this directory, "
                                             "mirroring the layout of directory inputs")
    out.add_argument("-i", "--in-place", action="store_true", help="overwrite the input files")
    p.add_argument("-k", "--key", help="secret key for the mapping; visible in the process "
                                       f"list, so prefer --key-file or ${KEY_ENV}")
    p.add_argument("--key-file", help="read the secret key from this file")
    p.add_argument("--save-key", help="write the key that was used to this file (mode 600)")
    p.add_argument("-r", "--reverse", action="store_true",
                   help="restore original addresses (needs the redaction key and the same "
                        "--v6-bits / exclude settings)")
    p.add_argument("-x", "--exclude", action="append", default=[], metavar="CIDR",
                   help="leave addresses in this IPv4 or IPv6 network untouched (repeatable)")
    p.add_argument("--no-default-excludes", action="store_true",
                   help="also redact " + ", ".join(DEFAULT_EXCLUDES))
    p.add_argument("--no-dashed", action="store_true",
                   help="don't redact dashed hostnames like ip-10-1-2-3")
    p.add_argument("--v6-bits", type=int, default=48, choices=V6_BITS_CHOICES,
                   help="how many leading IPv6 bits to rewrite (default 48: the site prefix)")
    p.add_argument("--map-out", metavar="CSV",
                   help="write an original,redacted CSV of every address seen "
                        "(mode 600; sensitive: contains the real addresses)")
    p.add_argument("-v", "--verbose", action="store_true", help="print a summary to stderr")
    return p


def main(argv=None):
    args = build_parser().parse_args(argv)
    inputs = collect_inputs(args.files)
    if args.output and (len(inputs) != 1 or any(os.path.isdir(f) for f in args.files)):
        sys.exit("error: --output takes a single input file; use --out-dir for several")
    if (args.in_place or args.out_dir) and "-" in args.files:
        sys.exit("error: --in-place/--out-dir need file arguments, not stdin")
    if not (args.in_place or args.out_dir) and any(os.path.isdir(f) for f in args.files):
        sys.exit("error: directory inputs need --out-dir or --in-place")
    for net in args.exclude:
        try:
            ipaddress.ip_network(net, strict=False)
        except ValueError:
            sys.exit(f"error: bad --exclude network: {net}")

    key, generated = resolve_key(args)
    if not generated and len(key) < MIN_KEY_LEN:
        warn(f"warning: key is shorter than {MIN_KEY_LEN} characters; anyone holding one "
             f"original/redacted pair could brute-force it and undo the redaction")
    excludes = list(args.exclude)
    if not args.no_default_excludes:
        excludes += DEFAULT_EXCLUDES
    redactor = Redactor(key, excludes, dashed=not args.no_dashed, v6_bits=args.v6_bits)

    failed = done = 0
    for src, dst in plan_outputs(inputs, args):
        try:
            if src.lower().endswith(XLSX_EXTS):
                if dst == "-":
                    raise SkipFile("spreadsheet output needs -o, -d or -i")
                data = redact_xlsx(redactor, src, args.reverse)
            else:
                with (sys.stdin.buffer if src == "-" else open(src, "rb")) as f:
                    raw = f.read()
                data = redact_bytes(redactor, raw, args.reverse)
            if dst != "-":
                os.makedirs(os.path.dirname(os.path.abspath(dst)), exist_ok=True)
            write_bytes(dst, data)
            done += 1
        except (SkipFile, OSError, UnicodeError) as e:
            failed += 1
            warn(f"skipped {src}: {e}")
        except Exception as e:  # e.g. a corrupt workbook; keep going with the rest
            failed += 1
            warn(f"skipped {src}: {type(e).__name__}: {e}")

    if args.save_key:
        write_bytes(args.save_key, (key + "\n").encode(), mode=0o600)
    if generated and not args.save_key:
        warn(f"generated key {key} - pass it with --key-file/--key to keep the "
             f"same mapping across runs or to --reverse")
    if args.map_out:
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(["original", "redacted"])
        for orig, red in sorted(redactor.seen.items(), key=lambda kv: (kv[0].version, kv[0])):
            w.writerow([orig, red])
        write_bytes(args.map_out, buf.getvalue().encode(), mode=0o600)
    if redactor.collisions:
        warn(f"warning: {len(redactor.collisions)} redacted address(es) landed in an "
             f"excluded range ({', '.join(map(str, sorted(redactor.collisions, key=lambda a: (a.version, a))))}); "
             f"--reverse will leave those as-is")
    if args.verbose:
        action = "restored" if args.reverse else "redacted"
        warn(f"{action} {redactor.count} address(es) in {done} file(s)"
             + (f", {len(redactor.seen)} unique" if not args.reverse else "")
             + (f"; {failed} file(s) skipped" if failed else ""))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
